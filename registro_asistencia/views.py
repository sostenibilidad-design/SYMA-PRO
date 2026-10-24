import json
import logging
import openpyxl
import calendar
import time

from decimal import Decimal
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from django.shortcuts import render, get_object_or_404, redirect
from django.db.models import Q
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse, Http404
from datetime import date, datetime, timedelta
from django.utils import timezone
from django.db.models import Count
from django.db import transaction
from django.urls import reverse
from django.utils.http import url_has_allowed_host_and_scheme
from django.core.validators import URLValidator
from django.core.exceptions import ValidationError
from django.views.decorators.http import require_POST
from django.utils.text import get_valid_filename

from .models import MaestroAsistencia,RegistroEntradaSalida,RegistroComida,EntradaSalidaMultiple,HoraHistorica,Festivo
from .festivos import FestivosNoDisponibles
from .liquidacion import (
    jornada_min, JORNADAS_POR_DEFECTO, DiaEmpleado, dias_habiles, liquidar, extras_del_dia, horas_de, comida_habitual_min,
    comidas_habituales_del_turno, minutos_efectivos, es_habil, ultimo_dia_habil_antes,
    HEO, NOCTURNA, DOM_DIURNA, DOM_NOCTURNA, ASISTENCIA, RETIRO, DESCANSO, TIPOS, NOMBRES_TIPO,
)
from proyectos.models import Proyecto
from core.excel import texto_excel
from usuario.decoradores import permiso_requerido
from personal.google_drive_utils import fetch_and_store_empleados
from personal.tasks import root_folder_id
from personal.models import Empleado
from medicion_rendimiento.models import MedicionCuadrilla, Cumplimiento

logger = logging.getLogger(__name__)


def permiso_registro(accion, json=False):
    """Permiso de este módulo (app y submódulo `registro_asistencia`; ver scan_funcionalidades.py): ver, registrar, actualizar o descargar."""
    return permiso_requerido('registro_asistencia', 'registro_asistencia', accion, json)


@permiso_registro('registrar', json=True)
@require_POST
def sincronizar_drive_inicial(request, id_proyecto):
    try:
        fetch_and_store_empleados(root_folder_id=root_folder_id, modulo='registro')
    except Exception:
        # El detalle va al registro del servidor; al navegador no se le muestran rutas ni mensajes internos
        logger.exception("Falló la sincronización con Drive")
        return JsonResponse({'success': False, 'error': 'No se pudo sincronizar con Drive. Intenta de nuevo.'}, status=500)

    # Pequeña pausa para que la animación del botón se vea fluida
    time.sleep(1)
    return JsonResponse({'success': True, 'mensaje': 'Base de datos sincronizada'})


@permiso_registro('ver')
def selector_asistencia(request):
    proyectos = Proyecto.objects.all()
    context = {
        'proyectos': proyectos,
        'url_destino': 'registro_anios',
    }
    return render(request, 'registro_asistencia/selector_asistencia.html', context)

MESES = [
    'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
    'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre',
]


def rango_quincena(anio, mes, quincena):
    """Devuelve (fecha_inicio, fecha_fin) de la quincena 1 (días 1-15) o 2 (16-fin de mes)."""
    if quincena == 1:
        return date(anio, mes, 1), date(anio, mes, 15)
    ultimo_dia = calendar.monthrange(anio, mes)[1]
    return date(anio, mes, 16), date(anio, mes, ultimo_dia)


def fechas_registradas(proyecto):
    """Fechas (ISO) del proyecto que ya tienen un registro guardado: no se puede registrar la misma fecha dos veces."""
    return [
        f.isoformat() for f in MaestroAsistencia.objects.filter(proyecto=proyecto, estado='Completado')
        .order_by('fecha').values_list('fecha', flat=True)
    ]


def festivos_para_calendario():
    """Festivos (ISO) del año anterior, el actual y el siguiente, de la copia guardada en la base de datos (ver festivos.py)."""
    anio = date.today().year
    return [
        f.isoformat() for f in Festivo.objects.filter(fecha__year__range=(anio - 1, anio + 1))
        .order_by('fecha').values_list('fecha', flat=True)
    ]


def limpiar_borradores_vencidos():
    """Elimina los maestros en Borrador con más de 24 horas sin guardarse."""
    limite = timezone.now() - timedelta(hours=24)
    MaestroAsistencia.objects.filter(estado='Borrador', creado_en__lt=limite).delete()


def texto_dias(cantidad):
    return f"{cantidad} día registrado" if cantidad == 1 else f"{cantidad} días registrados"


DIAS_SEMANA = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']


def construir_migas(proyecto, anio=None, mes=None, quincena=None, dia=None):
    """Proyecto > Año > Mes > Quincena > Día, cada miga enlaza a su nivel."""
    pk = proyecto.id_proyectos
    migas = [{'nombre': proyecto.nombre, 'url': reverse('registro_anios', args=[pk])}]
    if anio:
        migas.append({'nombre': str(anio), 'url': reverse('registro_meses', args=[pk, anio])})
    if mes:
        migas.append({'nombre': MESES[mes - 1], 'url': reverse('registro_quincenas', args=[pk, anio, mes])})
    if quincena:
        migas.append({'nombre': f"Q{quincena}", 'url': reverse('registro_dias', args=[pk, anio, mes, quincena])})
    if dia:
        migas.append({'nombre': str(dia), 'url': reverse('registro_asistencia', args=[pk, anio, mes, quincena, dia])})
    return migas


def horas_minimas(inicio, fin):
    """Horas mínimas (base) de una quincena: la jornada de todos sus días hábiles (8,5 h de lunes a jueves y 8 h el viernes;
    los fines de semana y festivos no cuentan). No depende de los empleados ni de lo registrado: es lo que se espera trabajar.
    None si no se pudo consultar el calendario de festivos."""
    jornadas = jornadas_por_dia()
    try:
        return horas_de(sum(jornada_min(d, jornadas) for d in dias_habiles(inicio, fin)))
    except FestivosNoDisponibles:
        return None


@permiso_registro('ver')
def registro_niveles(request, id_proyecto, anio=None, mes=None, quincena=None):
    """Vistas 1 (años), 2 (meses), 3 (quincenas) y 4 (días) del registro de entrada y salida."""
    proyecto_actual = get_object_or_404(Proyecto, pk=id_proyecto)
    if mes is not None and not 1 <= mes <= 12:
        raise Http404("Mes inválido")
    if quincena is not None and quincena not in (1, 2):
        raise Http404("Quincena inválida")

    hoy = date.today()
    limpiar_borradores_vencidos()
    # Solo cuentan los días guardados; los borradores no se muestran.
    maestros = MaestroAsistencia.objects.filter(proyecto=proyecto_actual, estado='Completado')

    context = {
        'proyecto_actual': proyecto_actual,
        'migas': construir_migas(proyecto_actual, anio, mes, quincena),
        'fechas_registradas': fechas_registradas(proyecto_actual),
        'festivos_calendario': festivos_para_calendario(),
    }

    if anio is None:
        # Vista 1: años con registros (el año actual siempre aparece para poder registrar)
        conteo = dict(
            maestros.values_list('fecha__year').annotate(total=Count('fecha__month', distinct=True)).values_list('fecha__year', 'total')
        )
        conteo.setdefault(hoy.year, 0)
        context.update({
            'nivel': 1,
            'carpetas': [
                {
                    'nombre': str(a),
                    'detalle': f"{conteo[a]} mes registrado" if conteo[a] == 1 else f"{conteo[a]} meses registrados",
                    'actual': a == hoy.year,
                    'url': reverse('registro_meses', args=[id_proyecto, a]),
                }
                for a in sorted(conteo, reverse=True)
            ],
            'url_volver': reverse('selector_asistencia'),
        })

    elif mes is None:
        # Vista 2: meses del año con registros (el mes actual siempre aparece si es el año en curso)
        conteo = dict(
            maestros.filter(fecha__year=anio)
            .values_list('fecha__month').annotate(total=Count('id')).values_list('fecha__month', 'total')
        )
        if anio == hoy.year:
            conteo.setdefault(hoy.month, 0)
        context.update({
            'nivel': 2,
            'carpetas': [
                {
                    'nombre': MESES[m - 1],
                    'detalle': texto_dias(conteo[m]),
                    'actual': anio == hoy.year and m == hoy.month,
                    'url': reverse('registro_quincenas', args=[id_proyecto, anio, m]),
                }
                for m in sorted(conteo, reverse=True)
            ],
            'url_volver': reverse('registro_anios', args=[id_proyecto]),
        })

    elif quincena is None:
        # Vista 3: las dos quincenas del mes, con las horas mínimas (base) de cada una
        quincenas = []
        for numero in (1, 2):
            inicio, fin = rango_quincena(anio, mes, numero)
            quincenas.append({
                'codigo': f"Q{numero}",
                'horas_base': horas_minimas(inicio, fin),
                'rango': f"{inicio.day} al {fin.day} de {MESES[mes - 1].lower()}",
                'estado': 'cerrada' if fin < hoy else 'abierta',
                'url': reverse('registro_dias', args=[id_proyecto, anio, mes, numero]),
            })
        context.update({
            'nivel': 3,
            'quincenas': quincenas,
            'url_volver': reverse('registro_meses', args=[id_proyecto, anio]),
        })

    else:
        # Vista 4: una carpeta por cada día de la quincena; solo las de días guardados se pueden abrir
        inicio, fin = rango_quincena(anio, mes, quincena)
        guardados = set(maestros.filter(fecha__range=(inicio, fin)).values_list('fecha', flat=True))
        fechas = [inicio + timedelta(days=n) for n in range((fin - inicio).days + 1)]
        context.update({
            'nivel': 4,
            'dias': [
                {
                    'fecha_iso': f.isoformat(),
                    'numero': f.day,
                    'dia_semana': DIAS_SEMANA[f.weekday()],
                    'mes': MESES[f.month - 1][:3],
                    'hoy': f == hoy,
                    'registrado': f in guardados,
                    'festivo': False,
                    'url': reverse('registro_asistencia', args=[id_proyecto, anio, mes, quincena, f.day]) if f in guardados else '',
                }
                for f in fechas
            ],
            'url_volver': reverse('registro_quincenas', args=[id_proyecto, anio, mes]),
            'url_eliminar_dias': reverse('eliminar_dias', args=[id_proyecto]),
            'url_descarga': reverse('descargar_dias_excel', args=[id_proyecto, anio, mes, quincena]),
            'url_informe_validar': reverse('informe_quincenal_validar', args=[id_proyecto, anio, mes, quincena]),
            'url_resumen_validar': reverse('resumen_actividades_validar', args=[id_proyecto, anio, mes, quincena]),
        })

    context['salto_config'] = construir_salto_config(context, maestros, id_proyecto, anio, hoy)
    return render(request, 'registro_asistencia/registro_niveles.html', context)


def construir_salto_config(context, maestros, id_proyecto, anio, hoy):
    """Datos del salto rápido (buscador de las vistas 1 a 4) que lee salto_rapido.js."""
    base = reverse('registro_anios', args=[id_proyecto])
    anios = set(maestros.values_list('fecha__year', flat=True))
    anios.add(hoy.year)

    if context['nivel'] == 3:
        opciones = [{'texto': q['codigo'], 'detalle': q['rango'], 'url': q['url']} for q in context['quincenas']]
    elif context['nivel'] == 4:
        opciones = [
            {'texto': f"{d['dia_semana']} {d['numero']}", 'detalle': d['fecha_iso'], 'url': d['url']}
            for d in context['dias'] if d['registrado']
        ]
    else:
        opciones = [{'texto': c['nombre'], 'detalle': c['detalle'], 'url': c['url']} for c in context['carpetas']]

    return {
        'hoy': hoy.isoformat(),
        'anio_contexto': anio or hoy.year,
        'anios': sorted(anios),
        'rutas': {
            'anio': f"{base}{{anio}}/",
            'mes': f"{base}{{anio}}/{{mes}}/",
            'quincena': f"{base}{{anio}}/{{mes}}/{{q}}/",
            'dia': f"{base}{{anio}}/{{mes}}/{{q}}/{{dia}}/",
        },
        'opciones': opciones,
    }


@permiso_registro('ver')
def registro_asistencia(request, id_proyecto, anio, mes, quincena, dia):
    """Vista 5: tabla de registros del día seleccionado."""
    proyecto_actual = get_object_or_404(Proyecto, pk=id_proyecto)
    if not 1 <= mes <= 12 or quincena not in (1, 2):
        raise Http404("Quincena inválida")
    try:
        fecha = date(anio, mes, dia)
    except ValueError:
        raise Http404("Día inválido")
    if (1 if fecha.day <= 15 else 2) != quincena:
        raise Http404("El día no pertenece a la quincena")

    registros_db = RegistroEntradaSalida.objects.filter(
        maestro__proyecto=proyecto_actual,
        maestro__fecha=fecha,
    ).select_related('empleado', 'maestro').prefetch_related(
        'comidas', 'horas_multiples', 'horas_historicas',
    ).order_by('empleado__nombre_completo')

    # Horas extra del día por tipo (lo que pasa de la jornada, o todo el turno en sábado, domingo y festivo)
    registros_db = list(registros_db)
    jornadas = jornadas_por_dia()
    for r in registros_db:
        r.extras_dia, r.faltan_dia = [], None
        if r.novedad != ASISTENCIA:
            continue
        try:
            extras, _, falta = extras_del_dia(dia_de_registro(r), jornadas)
        except FestivosNoDisponibles:
            r.extras_dia = None            # sin calendario de festivos no se puede clasificar el día
            continue
        r.extras_dia = [(NOMBRES_TIPO[t], horas_de(extras[t])) for t in TIPOS if extras[t] > 0]
        r.faltan_dia = horas_de(falta) if falta > 0 else None

    context = {
        'proyecto_actual': proyecto_actual,
        'registros': registros_db,
        'fechas_registradas': fechas_registradas(proyecto_actual),
        'festivos_calendario': festivos_para_calendario(),
        'ya_registrado': request.GET.get('ya_registrado') == '1',
        'fecha_dia': fecha,
        'migas': construir_migas(proyecto_actual, anio, mes, quincena, dia),
        'url_volver': reverse('registro_dias', args=[id_proyecto, anio, mes, quincena]),
    }
    return render(request, 'registro_asistencia/registro_asistencia.html', context)


def fecha_valida(texto):
    """True si `texto` es una fecha AAAA-MM-DD real (la fecha del registro llega por la URL o la sesión)."""
    try:
        datetime.strptime(texto, "%Y-%m-%d")
    except (TypeError, ValueError):
        return False
    return True


def url_quincena_de_fecha(id_proyecto, fecha_str):
    """URL de la vista 5 (tabla del día) de la fecha dada (o la vista de años si la fecha no es válida)."""
    try:
        fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return reverse('registro_anios', args=[id_proyecto])
    quincena = 1 if fecha.day <= 15 else 2
    return reverse('registro_asistencia', args=[id_proyecto, fecha.year, fecha.month, quincena, fecha.day])


def empleados_retirados_antes_de(fecha):
    """Cédulas de los empleados con novedad «Retiro» registrada en un día anterior a `fecha`."""
    return RegistroEntradaSalida.objects.filter(
        novedad=RETIRO, maestro__estado='Completado', maestro__fecha__lt=fecha,
    ).values('empleado_id')


@permiso_registro('registrar')
def paso_1_novedades(request, id_proyecto):
    proyecto_actual = get_object_or_404(Proyecto, pk=id_proyecto)
    fecha_str = request.GET.get('fecha')

    if not fecha_valida(fecha_str):
        return redirect('registro_anios', id_proyecto=id_proyecto)

    # Excepción: una fecha solo se registra una vez; si ya está guardada, se muestra su tabla
    if MaestroAsistencia.objects.filter(proyecto=proyecto_actual, fecha=fecha_str, estado='Completado').exists():
        return redirect(f"{url_quincena_de_fecha(id_proyecto, fecha_str)}?ya_registrado=1")

    if request.method == 'POST':
        # Leemos el array de IDs ocultos que creó JavaScript
        empleados_confirmados = request.POST.getlist('empleados_confirmados[]')
        
        datos_paso1 = {}
        
        # Solo empleados de este proyecto y novedades conocidas: lo demás que llegue en el POST se ignora
        ids_del_proyecto = {str(pk) for pk in Empleado.objects.filter(ubicacion=proyecto_actual).values_list('pk', flat=True)}
        for emp_id in empleados_confirmados:
            novedad = request.POST.get(f'novedad_{emp_id}')
            if emp_id in ids_del_proyecto and novedad in MAPEO_DB_NOVEDADES:
                datos_paso1[emp_id] = novedad
                
        # Guardamos en la mochila (Sesión)
        request.session['datos_paso1'] = datos_paso1
        request.session['fecha_asistencia'] = fecha_str # Guardamos la fecha también
        
        # Saltamos al Paso 2 llevando la fecha en la URL
        url_paso2 = reverse('paso_2_hora_inicio', args=[id_proyecto])
        return redirect(f"{url_paso2}?fecha={fecha_str}")
        
    limpiar_borradores_vencidos()

    # Página desde la que se pulsó «Registrar»: la X del registro vuelve allí (si no llega, vuelve a las carpetas)
    origen = request.GET.get('origen', '')
    if origen.startswith('/') and url_has_allowed_host_and_scheme(origen, allowed_hosts={request.get_host()}):
        request.session['origen_registro'] = origen
    else:
        request.session.pop('origen_registro', None)

    # get_or_create busca si ya existe un registro para ese proyecto y esa fecha.
    # Si ya existe, simplemente lo carga.
    maestro_actual, creado = MaestroAsistencia.objects.get_or_create(
        proyecto=proyecto_actual,
        fecha=fecha_str,
        defaults={
            'estado': 'Borrador',
            'usuario_creador': request.user
        }
    )

    cargos_permitidos = Empleado.objects.filter(
        Q(cargo__icontains='ayudante') |
        Q(cargo__icontains='oficial') |
        Q(cargo__icontains='maestro')|
        Q(cargo__icontains='almacenista'),
        ubicacion=proyecto_actual
    ).exclude(
        # Retirado: inactivo desde el día siguiente a la fecha de su novedad «Retiro»
        pk__in=empleados_retirados_antes_de(fecha_str)
    ).order_by('nombre_completo')

    context = {
        'proyecto_actual': proyecto_actual,
        'maestro': maestro_actual,
        'fecha_seleccionada': fecha_str,
        'empleados': cargos_permitidos,
        'paso_actual': 1,
        'total_pasos': 5,
        'titulo_paso': 'Asignar una novedad a cada empleado',
        'titulo_caja_derecha': 'Asignar novedad',
        'url_siguiente': 'paso_2_hora_inicio',
        'asignaciones_previas': request.session.get('datos_paso1', {}),   # se rearman al volver con «Atrás»
    }
    return render(request, 'registro_asistencia/novedades.html', context)

def descartar_multiples_si_cambia(session, clave, nuevas):
    """Quien cambia su hora de entrada o de salida pierde sus entradas y salidas múltiples (evita inconsistencias).

    `clave` es 'datos_paso2' (entradas) o 'datos_paso3' (salidas) y `nuevas` las horas que se van a guardar.
    """
    anteriores = session.get(clave, {})
    cambiaron = {emp for emp, hora in nuevas.items() if emp in anteriores and anteriores[emp] != hora}
    multiples = session.get('datos_multiples', {})
    if cambiaron & multiples.keys():
        session['datos_multiples'] = {emp: filas for emp, filas in multiples.items() if emp not in cambiaron}


def empleados_con_multiples(session):
    """IDs de los empleados que ya tienen entradas y salidas múltiples en el registro en curso."""
    return [emp for emp, filas in session.get('datos_multiples', {}).items() if filas]


@permiso_registro('registrar')
def paso_2_hora_inicio(request, id_proyecto):
    proyecto_actual = get_object_or_404(Proyecto, pk=id_proyecto)
    fecha_str = request.GET.get('fecha') or request.session.get('fecha_asistencia')
    if not fecha_valida(fecha_str):
        return redirect('registro_anios', id_proyecto=id_proyecto)

    # ABRIMOS LA MOCHILA
    datos_paso1 = request.session.get('datos_paso1', {})

    if not datos_paso1:
        # Si intentan entrar directo a la URL, los devolvemos al inicio
        url_paso1 = reverse('paso_1_novedades', args=[id_proyecto])
        return redirect(f"{url_paso1}?fecha={fecha_str}")

    if request.method == 'POST':
        
        datos_paso2 = {}
        
        # Iteramos solo sobre los que tuvieron asistencia normal
        for emp_id, novedad in datos_paso1.items():
            if novedad == 'asistencia_normal':
                # Buscamos la hora exacta que el JavaScript escondió en los chips
                hora_asignada = request.POST.get(f'hora_inicio_{emp_id}')
                if hora_asignada and convertir_a_24h(hora_asignada):
                    datos_paso2[emp_id] = hora_asignada
                    
        # Guardamos las horas en la mochila
        descartar_multiples_si_cambia(request.session, 'datos_paso2', datos_paso2)
        request.session['datos_paso2'] = datos_paso2
        
        # Saltamos al Paso 3 llevando la fecha
        url_paso3 = reverse('paso_3_hora_salida', args=[id_proyecto])
        return redirect(f"{url_paso3}?fecha={fecha_str}")
    
    # Excluir Grupos A y B
    ids_asistencia_normal = [
        emp_id for emp_id, novedad in datos_paso1.items() 
        if novedad == 'asistencia_normal'
    ]
    
    # Buscar en la base de datos SÓLO a los seleccionados
    empleados_activos = Empleado.objects.filter(pk__in=ids_asistencia_normal).order_by('nombre_completo')
    
    context = {
        'proyecto_actual': proyecto_actual,
        'fecha_seleccionada': fecha_str,
        'empleados': empleados_activos,
        'paso_actual': 2,
        'total_pasos': 5,
        'titulo_paso': 'Asignar hora de inicio a cada empleado',
        'titulo_caja_derecha': 'Asignar hora de inicio (masivo)',
        'url_siguiente': 'paso_3_hora_salida',
        'asignaciones_previas': request.session.get('datos_paso2', {}),
        'con_multiples': empleados_con_multiples(request.session),
        'url_atras': f"{reverse('paso_1_novedades', args=[id_proyecto])}?fecha={fecha_str}",
    }
    return render(request, 'registro_asistencia/hora_de_entrada.html', context)

@permiso_registro('registrar')
def paso_3_hora_salida(request, id_proyecto):
    proyecto_actual = get_object_or_404(Proyecto, pk=id_proyecto)
    fecha_str = request.GET.get('fecha') or request.session.get('fecha_asistencia')
    if not fecha_valida(fecha_str):
        return redirect('registro_anios', id_proyecto=id_proyecto)

    # 1. Abrimos la mochila para ver quiénes traen hora de entrada del Paso 2
    datos_paso2 = request.session.get('datos_paso2', {})

    if not datos_paso2:
        # Si saltan directo aquí sin pasar por el 2, los devolvemos
        url_paso2 = reverse('paso_2_hora_inicio', args=[id_proyecto])
        return redirect(f"{url_paso2}?fecha={fecha_str}")

    # 2. PROCESAR EL FORMULARIO CUANDO LE DAN A "SIGUIENTE"
    if request.method == 'POST':
        datos_paso3 = {}
        
        # Iteramos sobre los que traen hora de inicio
        for emp_id in datos_paso2.keys():
            # Buscamos la hora de salida que el JavaScript escondió
            hora_salida = request.POST.get(f'hora_salida_{emp_id}')
            if hora_salida and convertir_a_24h(hora_salida):
                datos_paso3[emp_id] = hora_salida
                
        # Guardamos en la mochila
        descartar_multiples_si_cambia(request.session, 'datos_paso3', datos_paso3)
        request.session['datos_paso3'] = datos_paso3
        
        # Saltamos al Paso 4
        url_paso4 = reverse('paso_4_actividades', args=[id_proyecto])
        return redirect(f"{url_paso4}?fecha={fecha_str}")

    # 3. MOSTRAR LA PÁGINA (MÉTODO GET)
    # Extraemos solo los IDs (las llaves del diccionario de la mochila)
    ids_con_entrada = list(datos_paso2.keys())
    
    # Buscamos a esos empleados para mandarlos al panel izquierdo
    empleados_activos = Empleado.objects.filter(pk__in=ids_con_entrada).order_by('nombre_completo')
    
    context = {
        'proyecto_actual': proyecto_actual,
        'fecha_seleccionada': fecha_str,
        'empleados': empleados_activos, # <-- ¡Esto hará que aparezcan en la lista!
        'paso_actual': 3,
        'total_pasos': 5,
        'titulo_paso': 'Asignar hora de salida a cada empleado',
        'titulo_caja_derecha': 'Asignar hora de salida (masivo)',
        'url_siguiente': 'paso_4_actividades',
        'asignaciones_previas': request.session.get('datos_paso3', {}),
        'con_multiples': empleados_con_multiples(request.session),
        'url_atras': f"{reverse('paso_2_hora_inicio', args=[id_proyecto])}?fecha={fecha_str}",
    }
    return render(request, 'registro_asistencia/hora_de_salida.html', context)

@permiso_registro('registrar')
def paso_4_actividades(request, id_proyecto):
    proyecto_actual = get_object_or_404(Proyecto, pk=id_proyecto)
    fecha_str = request.GET.get('fecha') or request.session.get('fecha_asistencia')
    if not fecha_valida(fecha_str):
        return redirect('registro_anios', id_proyecto=id_proyecto)

    # 1. Abrimos la mochila: Solo pasan los que tienen hora de salida
    datos_paso3 = request.session.get('datos_paso3', {})
    datos_paso4 = request.session.get('datos_paso4', {})

    if not datos_paso3:
        url_paso3 = reverse('paso_3_hora_salida', args=[id_proyecto])
        return redirect(f"{url_paso3}?fecha={fecha_str}")

    # 2. PROCESAR EL FORMULARIO CUANDO LE DAN A "SIGUIENTE"
    if request.method == 'POST':
        datos_paso4 = {}
        # Guardamos la actividad asignada a cada empleado
        for emp_id in datos_paso3.keys():
            actividad_asignada = request.POST.get(f'actividad_{emp_id}')
            if actividad_asignada and len(actividad_asignada) <= 255:
                datos_paso4[emp_id] = actividad_asignada
                
        request.session['datos_paso4'] = datos_paso4
        url_paso5 = reverse('paso_5_vista_previa', args=[id_proyecto])
        return redirect(f"{url_paso5}?fecha={fecha_str}")

    # 3. MOSTRAR LA PÁGINA (MÉTODO GET)
    ids_con_salida = list(datos_paso3.keys())
    
    empleados_activos = list(Empleado.objects.filter(pk__in=ids_con_salida).order_by('nombre_completo'))
    
    # Extraer las actividades disponibles para el proyecto actual
    actividades_disponibles = list(Cumplimiento.objects.filter(proyecto=proyecto_actual).values_list('actividad', flat=True).distinct())

    # Escanear coincidencias en MedicionCuadrilla y "pegarle" el dato al empleado
    for emp in empleados_activos:
        medicion = MedicionCuadrilla.objects.filter(
            fecha=fecha_str, 
            empleados=emp,
            proyecto=proyecto_actual
        ).first()
        
        # El .strip() es VITAL para quitar espacios invisibles que rompen el HTML
        emp.actividad_sugerida = medicion.actividad.strip() if medicion and medicion.actividad else ""
        
        # Si ya se había elegido una actividad y se volvió con «Atrás», se conserva
        emp.actividad_sugerida = datos_paso4.get(str(emp.pk)) or emp.actividad_sugerida

    context = {
        'proyecto_actual': proyecto_actual,
        'fecha_seleccionada': fecha_str,
        'empleados': empleados_activos, 
        'actividades_disponibles': actividades_disponibles,
        'paso_actual': 4,
        'total_pasos': 5,
        'titulo_paso': 'Asignación y confirmación de las actividades',
        'url_siguiente': 'paso_5_vista_previa',
        'url_atras': f"{reverse('paso_3_hora_salida', args=[id_proyecto])}?fecha={fecha_str}",
    }
    return render(request, 'registro_asistencia/actividades.html', context)

def convertir_a_24h(hora_texto):
    """Traduce el texto para que el reloj HTML5 lo entienda internamente"""
    if not hora_texto or hora_texto == '--:--':
        return ''
    try:
        texto_limpio = hora_texto.replace('.', '').strip().upper()
        dt = datetime.strptime(texto_limpio, "%I:%M %p")
        return dt.strftime("%H:%M")
    except Exception:
        return ''
    
def hora_o_none(texto):
    """'07:00' (o '07:00:00') -> time; None si está vacío o no se puede leer."""
    try:
        return datetime.strptime(texto[:5], "%H:%M").time()
    except (ValueError, TypeError):
        return None


TIPOS_COMIDA = ('desayuno', 'almuerzo', 'cena')

def jornadas_por_dia():
    """Jornada en minutos por día de la semana (0 = lunes). Se cambia en JORNADAS_POR_DEFECTO (liquidacion.py)."""
    return dict(JORNADAS_POR_DEFECTO)


def horas_comida_habituales(proyecto, fecha, entrada, salida):
    """Horas de comida que se restan por defecto: solo las franjas del proyecto que el turno cruza (ver COMIDAS_POR_PROYECTO en liquidacion.py).

    Guayabal y los demás: almuerzo de 12:00 a 13:00. Ciudadela Andina: desayuno de 8:00 a 8:30 y almuerzo de 12:00 a 13:00.
    `entrada` y `salida` son time o None; las excepciones de la herramienta de comidas reemplazan este valor.
    """
    return (Decimal(comida_habitual_min(proyecto.nombre, fecha, entrada, salida)) / 60).quantize(Decimal('0.01'))


def horas_comida_registro(proyecto, fecha, entrada, salida, excepciones):
    """Horas de comida de un registro nuevo: las excepciones [(inicio, fin)] reemplazan a la comida habitual del proyecto."""
    minutos = sum((fin.hour * 60 + fin.minute) - (ini.hour * 60 + ini.minute) for ini, fin in excepciones)
    if minutos:
        return (Decimal(minutos) / 60).quantize(Decimal('0.01'))
    return horas_comida_habituales(proyecto, fecha, entrada, salida)


def horas_de_comidas(comidas):
    """Horas de comida de los horarios [(inicio, fin)] tal cual: lo que no aparece, no se tomó."""
    minutos = sum((fin.hour * 60 + fin.minute) - (ini.hour * 60 + ini.minute) for ini, fin in comidas)
    return (Decimal(minutos) / 60).quantize(Decimal('0.01'))


def horas_base_del_dia(fecha):
    """Jornada de ese día en horas (8,5 h de lunes a jueves, 8 h el viernes; 0 en fin de semana y festivos)."""
    jornadas = jornadas_por_dia()
    try:
        return Decimal(jornada_min(fecha, jornadas)) / 60
    except FestivosNoDisponibles:
        # Sin calendario no se puede saber si es festivo: se guarda la jornada del día de la semana
        return Decimal(jornadas.get(fecha.weekday(), 0)) / 60


# Este es el diccionario para visualizar en la tabla (Frontend)
NOMBRES_NOVEDADES = {
    'asistencia_normal': 'Asistencia normal', 'ausentismo': 'Ausentismo',
    'sancion': 'Sanción', 'incapacidad_general': 'Incapacidad general',
    'incapacidad_trabajo': 'Incapacidad trab.', 'permiso_no_remunerado': 'Permiso no rem.',
    'permiso_remunerado': 'Permiso rem.', 'licencia': 'Licencia',
    'vacaciones': 'Vacaciones', 'retirado': 'Retirado',
    'descanso': 'Descanso',
}

# Este es el diccionario para traducir al modelo de Base de Datos
MAPEO_DB_NOVEDADES = {
    'asistencia_normal': 'Asistencia', 'ausentismo': 'Ausentismo',
    'sancion': 'Sanción', 'incapacidad_general': 'Incapacidad general',
    'incapacidad_trabajo': 'Incapacidad laboral', 'permiso_no_remunerado': 'Permiso no remunerado',
    'permiso_remunerado': 'Permiso remunerado', 'licencia': 'Licencia',
    'vacaciones': 'Vacaciones', 'retirado': 'Retiro',
    'descanso': 'Descanso',
}


def parsear_rango_horas(texto):
    """'06:30-06:50' (24 h) -> (time, time). None si el formato no es válido o la segunda hora no es mayor que la primera."""
    try:
        inicio, fin = (datetime.strptime(parte, "%H:%M").time() for parte in texto.split('-'))
    except (ValueError, AttributeError):
        return None
    return (inicio, fin) if fin > inicio else None


def parsear_rango_libre(texto):
    """'23:00-01:00' (24 h) -> (time, time); a diferencia de parsear_rango_horas, el ingreso puede ser de la madrugada siguiente.
    None si el formato no es válido o las dos horas son iguales."""
    try:
        a, b = (datetime.strptime(parte, "%H:%M").time() for parte in texto.split('-'))
    except (ValueError, AttributeError):
        return None
    return (a, b) if a != b else None


def _desde_entrada(entrada, hora):
    """Minutos desde la entrada hasta `hora` (una hora anterior a la entrada es del día siguiente)."""
    return (hora.hour * 60 + hora.minute - entrada.hour * 60 - entrada.minute) % 1440


def multiple_dentro_del_turno(entrada, salida, salida_m, ingreso_m):
    """True si la salida y el ingreso múltiples van dentro del turno (también el que pasa de medianoche): después de la
    entrada, el ingreso después de la salida y antes de la salida principal."""
    if not (entrada and salida and salida_m and ingreso_m) or entrada == salida:
        return False
    fin, a, b = _desde_entrada(entrada, salida), _desde_entrada(entrada, salida_m), _desde_entrada(entrada, ingreso_m)
    return 0 < a < b < fin


def _cruzan(r1, r2):
    """True si dos (salida, ingreso) se pisan; un ingreso de madrugada se mide como del día siguiente."""
    def minutos(r):
        a, b = r[0].hour * 60 + r[0].minute, r[1].hour * 60 + r[1].minute
        return a, (b if b > a else b + 1440)
    (a1, b1), (a2, b2) = minutos(r1), minutos(r2)
    # Se prueba también con uno de los dos corrido un día, por si el otro cae después de medianoche
    return any(a1 < b2 + dia and a2 + dia < b1 for dia in (-1440, 0, 1440))


def multiples_solapadas(rangos):
    """True si dos salidas múltiples [(salida, ingreso)] se cruzan entre sí (se restarían dos veces del turno)."""
    return any(_cruzan(r1, r2) for n, r1 in enumerate(rangos) for r2 in rangos[n + 1:])


# Herramientas opcionales: clave de la URL -> título de la pantalla
HERRAMIENTAS_OPCIONALES = {
    'comidas': 'Gestión de comidas (herramienta opcional)',
    'multiples': 'Entradas y salidas múltiples (herramienta opcional)',
    'comentarios': 'Comentarios (herramienta opcional)',
    'historicas': 'Horas históricas (herramienta opcional)',
}
MAX_HORAS_HISTORICAS = 4


def texto_hora(hhmm):
    """'17:00' -> '05:00 PM'; si no se puede leer, devuelve el texto tal cual."""
    try:
        return datetime.strptime(hhmm, "%H:%M").strftime("%I:%M %p")
    except ValueError:
        return hhmm


def a_decimal(texto):
    try:
        return Decimal(str(texto).replace(',', '.'))
    except Exception:
        return None


@permiso_registro('registrar')
def herramienta_opcional(request, id_proyecto, herramienta):
    """Herramientas opcionales del registro: comidas, entradas/salidas múltiples, comentarios y horas históricas.

    Todo se guarda en la sesión y se consolida en la base de datos en el paso 5.
    """
    if herramienta not in HERRAMIENTAS_OPCIONALES:
        raise Http404("Herramienta no encontrada")
    proyecto_actual = get_object_or_404(Proyecto, pk=id_proyecto)
    fecha_str = request.GET.get('fecha') or request.session.get('fecha_asistencia')
    if not fecha_valida(fecha_str):
        return redirect('registro_anios', id_proyecto=id_proyecto)

    datos_paso1 = request.session.get('datos_paso1', {})
    datos_paso2 = request.session.get('datos_paso2', {})
    datos_paso3 = request.session.get('datos_paso3', {})

    # Cada herramienta exige haber llegado al paso que la desbloquea
    requisito = {'comidas': datos_paso1, 'comentarios': datos_paso2, 'historicas': datos_paso2, 'multiples': datos_paso3}
    if not fecha_str or not requisito[herramienta]:
        return redirect('registro_anios', id_proyecto=id_proyecto)

    try:
        paso = min(max(int(request.GET.get('paso', 2)), 2), 5)
    except ValueError:
        paso = 2

    # Solo se permite volver a una ruta de este mismo sitio
    volver = request.POST.get('volver') or request.GET.get('volver', '')
    if not url_has_allowed_host_and_scheme(volver, allowed_hosts={request.get_host()}):
        volver = f"{reverse('paso_2_hora_inicio', args=[id_proyecto])}?fecha={fecha_str}"

    ids_asistencia = [emp_id for emp_id, novedad in datos_paso1.items() if novedad == 'asistencia_normal']
    ids_multiples = [emp_id for emp_id in datos_paso3 if emp_id in ids_asistencia]
    ids_todos = list(datos_paso1.keys())
    errores = []
    comentarios_posteados = None

    # ---------- Guardar ----------
    if request.method == 'POST':
        if herramienta == 'comidas':
            datos = {}
            for clave, valor in request.POST.items():
                partes = clave.split('__')
                if len(partes) != 3 or partes[0] != 'comida':
                    continue
                _, emp_id, tipo = partes
                if emp_id in ids_asistencia and tipo in TIPOS_COMIDA and parsear_rango_horas(valor):
                    datos.setdefault(emp_id, {})[tipo] = valor
            request.session['datos_comidas'] = datos

        elif herramienta == 'multiples':
            datos = {}
            for clave, valor in request.POST.items():
                partes = clave.split('__')
                if len(partes) != 3 or partes[0] != 'multiple' or partes[1] not in ids_multiples or not partes[2].isdigit():
                    continue
                emp_id, orden = partes[1], int(partes[2])
                horas = parsear_rango_libre(valor)
                entrada = hora_o_none(convertir_a_24h(datos_paso2.get(emp_id, '')))
                salida = hora_o_none(convertir_a_24h(datos_paso3.get(emp_id, '')))
                # La salida y el ingreso múltiples van dentro del turno: después de la entrada y antes de la salida principal
                if horas and multiple_dentro_del_turno(entrada, salida, *horas):
                    datos.setdefault(emp_id, []).append((orden, valor))
            # Las que se cruzan con otra ya aceptada se descartan (se restarían dos veces)
            multiples = {}
            for emp_id, items in datos.items():
                aceptadas = []
                for _, valor in sorted(items):
                    horas = parsear_rango_libre(valor)
                    if not multiples_solapadas(aceptadas + [horas]):
                        aceptadas.append(horas)
                        multiples.setdefault(emp_id, []).append(valor)
            request.session['datos_multiples'] = multiples

        elif herramienta == 'comentarios':
            datos = {}
            validar_url = URLValidator()
            for emp_id in ids_todos:
                comentario = request.POST.get(f'comentario_{emp_id}', '').strip()
                link = request.POST.get(f'link_del_comentario_{emp_id}', '').strip()
                if link:
                    try:
                        validar_url(link)
                    except ValidationError:
                        nombre = Empleado.objects.get(pk=emp_id).nombre_completo
                        errores.append(f"El enlace de soporte de {nombre} no es una URL válida.")
                if comentario or link:
                    datos[emp_id] = {'comentario': comentario, 'link': link}
            if errores:
                comentarios_posteados = datos
            else:
                request.session['datos_comentarios'] = datos

        else:  # historicas
            datos = {'pendientes': {}, 'deducir': {}}
            tipos_validos = {valor for valor, _ in HoraHistorica.TIPO_HORAS}
            for emp_id in ids_todos:
                for prefijo, clave in (('pend', 'pendientes'), ('ded', 'deducir')):
                    tipos = request.POST.getlist(f'{prefijo}_tipo_{emp_id}')
                    cantidades = request.POST.getlist(f'{prefijo}_cantidad_{emp_id}')
                    filas = []
                    vistos = set()
                    for tipo, cantidad in zip(tipos, cantidades):
                        cantidad = a_decimal(cantidad)
                        # Un empleado no puede repetir el tipo de hora dentro de la misma pestaña
                        if tipo in tipos_validos and tipo not in vistos and cantidad is not None and 0 < cantidad <= 24:
                            vistos.add(tipo)
                            filas.append({'tipo': tipo, 'cantidad': str(cantidad.quantize(Decimal('0.01')))})
                    if filas:
                        datos[clave][emp_id] = filas[:MAX_HORAS_HISTORICAS]
            request.session['datos_historicas'] = datos

        if not errores:
            return redirect(volver)

    # ---------- Mostrar ----------
    if herramienta == 'multiples':
        ids_lista = ids_multiples
    elif herramienta in ('comentarios', 'historicas'):
        ids_lista = ids_todos          # los comentarios y las horas históricas aplican a todos, también a quienes no trabajaron
    else:
        ids_lista = ids_asistencia
    empleados = Empleado.objects.filter(pk__in=ids_lista).order_by('nombre_completo')

    context = {
        'proyecto_actual': proyecto_actual,
        'fecha_seleccionada': fecha_str,
        'empleados': empleados,
        'herramienta': herramienta,
        'es_herramienta': True,
        # Comentarios y horas históricas se guardan al escribir: no llevan caja de Confirmados
        'sin_confirmados': herramienta in ('comentarios', 'historicas'),
        'paso_actual': paso,
        'total_pasos': 5,
        'titulo_paso': HERRAMIENTAS_OPCIONALES[herramienta],
        'mensaje_footer': 'Empleados asignados',
        'url_volver': volver,
        'errores': errores,
    }

    if herramienta == 'comidas':
        context['comidas_guardadas'] = request.session.get('datos_comidas', {})

    elif herramienta == 'multiples':
        context['multiples_guardadas'] = request.session.get('datos_multiples', {})
        context['principales'] = {
            emp_id: {
                'entrada': convertir_a_24h(datos_paso2.get(emp_id, '')),
                'salida': convertir_a_24h(datos_paso3.get(emp_id, '')),
            }
            for emp_id in ids_multiples
        }

    elif herramienta == 'comentarios':
        context['comentarios_guardadas'] = (
            comentarios_posteados if comentarios_posteados is not None else request.session.get('datos_comentarios', {})
        )

    else:  # historicas
        context['tipos_horas'] = HoraHistorica.TIPO_HORAS
        context['historicas_guardadas'] = request.session.get('datos_historicas', {'pendientes': {}, 'deducir': {}})

    return render(request, 'registro_asistencia/pasos_opcionales.html', context)


@permiso_registro('registrar')
@transaction.atomic   # o se guarda todo el día o nada; el bloqueo del maestro evita guardados dobles a la vez
def paso_5_vista_previa(request, id_proyecto):
    proyecto_actual = get_object_or_404(Proyecto, pk=id_proyecto)
    fecha_str = request.GET.get('fecha') or request.session.get('fecha_asistencia')
    if not fecha_valida(fecha_str):
        return redirect('registro_anios', id_proyecto=id_proyecto)

    datos_paso1 = request.session.get('datos_paso1', {}) 
    datos_paso2 = request.session.get('datos_paso2', {})  
    datos_paso3 = request.session.get('datos_paso3', {})  
    datos_paso4 = request.session.get('datos_paso4', {})  
    datos_comidas = request.session.get('datos_comidas', {})
    datos_multiples = request.session.get('datos_multiples', {})
    datos_comentarios = request.session.get('datos_comentarios', {})
    datos_historicas = request.session.get('datos_historicas', {'pendientes': {}, 'deducir': {}})

    if not datos_paso1:
        url_paso1 = reverse('paso_1_novedades', args=[id_proyecto])
        return redirect(f"{url_paso1}?fecha={fecha_str}")

    # Extraer las actividades disponibles
    actividades_disponibles = list(Cumplimiento.objects.filter(
        Q(cumplimiento_presupuestal__gt=0) | Q(cumplimiento_programado__gt=0),
        proyecto=proyecto_actual
    ).values_list('actividad', flat=True).distinct())

    errores = []
    if request.method == 'POST':
        # Comentario y enlace de soporte se editan en la tabla: el enlace, si se llena, debe ser una URL válida
        validar_url = URLValidator()
        datos_comentarios = {}
        # Solo cuentan los empleados que pasaron por el paso 1 de este registro
        ids_enviados = [clave.split('_')[-1] for clave in request.POST
                        if clave.startswith('final_novedad_') and clave.split('_')[-1] in datos_paso1]
        nombres = {str(e.pk): e.nombre_completo for e in Empleado.objects.filter(pk__in=ids_enviados)}
        for emp_id in ids_enviados:
            nombre = nombres.get(emp_id, emp_id)
            if request.POST.get(f'final_novedad_{emp_id}') not in MAPEO_DB_NOVEDADES:
                errores.append(f"La novedad de {nombre} no es válida.")
                continue
            for campo, texto in (('entrada', 'hora de entrada'), ('salida', 'hora de salida')):
                valor = request.POST.get(f'final_{campo}_{emp_id}', '').strip()
                if valor and hora_o_none(valor) is None:
                    errores.append(f"La {texto} de {nombre} no es válida.")
            if len(request.POST.get(f'final_actividad_{emp_id}', '')) > 255:
                errores.append(f"La actividad de {nombre} es demasiado larga.")
            # Quien queda con asistencia normal necesita todos los datos de los pasos obligatorios
            if request.POST.get(f'final_novedad_{emp_id}') == 'asistencia_normal':
                faltan = [texto for campo, texto in (('entrada', 'la hora de entrada'), ('salida', 'la hora de salida'), ('actividad', 'la actividad'))
                          if not request.POST.get(f'final_{campo}_{emp_id}', '').strip()]
                if faltan:
                    errores.append(f"Faltan datos obligatorios de {nombre}: {', '.join(faltan)}.")
            comentario = request.POST.get(f'final_comentario_{emp_id}', '').strip()
            link = request.POST.get(f'final_link_{emp_id}', '').strip()
            if link:
                try:
                    validar_url(link)
                except ValidationError:
                    errores.append(f"El enlace de soporte de {nombre} no es una URL válida.")
            if comentario or link:
                datos_comentarios[emp_id] = {'comentario': comentario, 'link': link}
        if errores:
            # No se guarda nada: se conserva lo escrito y se vuelve a mostrar la tabla con el aviso
            request.session['datos_comentarios'] = datos_comentarios

    # PROCESO DE GUARDADO EN BASE DE DATOS
    if request.method == 'POST' and not errores:
        # 1. Recuperar el Maestro que creamos en el Paso 1
        maestro_actual = get_object_or_404(MaestroAsistencia.objects.select_for_update(), proyecto=proyecto_actual, fecha=fecha_str)
        if maestro_actual.estado == 'Completado':
            # Otro registro de esa fecha se guardó antes (por ejemplo, desde otra pestaña): no se sobrescribe
            for k in ['datos_paso1', 'datos_paso2', 'datos_paso3', 'datos_paso4', 'datos_comidas', 'datos_multiples', 'datos_comentarios', 'datos_historicas', 'fecha_asistencia']:
                request.session.pop(k, None)
            return redirect(f"{url_quincena_de_fecha(id_proyecto, fecha_str)}?ya_registrado=1")
        
        # 2. Jornada de ese día (8,5 h de lunes a jueves, 8 h el viernes; 0 en fin de semana y festivos)
        horas_base_actual = horas_base_del_dia(maestro_actual.fecha)

        # 3. Leer los datos enviados por el HTML
        for key, value in request.POST.items():
            if key.startswith('final_novedad_'):
                # Extraemos el ID del empleado
                emp_pk = key.split('_')[-1]
                if emp_pk not in datos_paso1:
                    continue
                
                # Leemos todos los campos correspondientes a este empleado
                novedad_codigo = request.POST.get(f'final_novedad_{emp_pk}')
                hora_entrada = request.POST.get(f'final_entrada_{emp_pk}')
                hora_salida = request.POST.get(f'final_salida_{emp_pk}')
                actividad = request.POST.get(f'final_actividad_{emp_pk}')

                # Limpieza de datos (Django necesita None en vez de '' para campos TimeField)
                h_inicio = hora_entrada if hora_entrada else None
                h_fin = hora_salida if hora_salida else None
                actividad_final = actividad if actividad and actividad != 'Otra actividad manual' else None
                
                # Traducimos la novedad al modelo exacto
                novedad_db = MAPEO_DB_NOVEDADES.get(novedad_codigo, 'Asistencia')

                # Buscamos al empleado
                try:
                    empleado_obj = Empleado.objects.get(pk=emp_pk)
                except Empleado.DoesNotExist:
                    continue

                # 4. Guardamos o actualizamos el registro del empleado
                registro, _ = RegistroEntradaSalida.objects.update_or_create(
                    maestro=maestro_actual,
                    empleado=empleado_obj,
                    defaults={
                        'novedad': novedad_db,
                        'hora_inicio': h_inicio,
                        'hora_fin': h_fin,
                        'horas_base': horas_base_actual,
                        'actividad': actividad_final,
                        'comentario': datos_comentarios.get(str(emp_pk), {}).get('comentario') or None,
                        'link_del_comentario': datos_comentarios.get(str(emp_pk), {}).get('link') or None,
                    }
                )

                # Entradas y salidas múltiples: solo con asistencia normal; se reemplazan las anteriores
                registro.horas_multiples.all().delete()
                if novedad_codigo == 'asistencia_normal':
                    entrada_t, salida_t = hora_o_none(h_inicio), hora_o_none(h_fin)
                    # Si se cambió la entrada o la salida, las múltiples se eliminan para evitar inconsistencias
                    cambio_horario = ((h_inicio or '')[:5] != convertir_a_24h(datos_paso2.get(str(emp_pk), ''))
                                        or (h_fin or '')[:5] != convertir_a_24h(datos_paso3.get(str(emp_pk), '')))
                    borrar_multiples = cambio_horario or request.POST.get(f'final_borrar_multiples_{emp_pk}') == '1'
                    guardadas = []
                    for rango in ([] if borrar_multiples else datos_multiples.get(str(emp_pk), [])):
                        horas = parsear_rango_libre(rango)
                        if horas and multiple_dentro_del_turno(entrada_t, salida_t, *horas) and not multiples_solapadas(guardadas + [horas]):
                            guardadas.append(horas)
                            EntradaSalidaMultiple.objects.create(
                                registro=registro,
                                hora_de_salida_multiple=horas[0],
                                hora_de_ingreso_multiple=horas[1],
                            )

                # Horas históricas: pendientes (a favor) y a deducir (en contra)
                registro.horas_historicas.all().delete()
                for clave, accion in (('pendientes', 'Pendiente'), ('deducir', 'Deducir')):
                    for fila in datos_historicas.get(clave, {}).get(str(emp_pk), []):
                        HoraHistorica.objects.create(
                            registro=registro,
                            tipo_de_horas=fila['tipo'],
                            cantidad_de_horas=Decimal(fila['cantidad']),
                            accion=accion,
                        )

                # Comidas (herramienta opcional): solo para asistencia normal; se reemplazan las anteriores.
                # Sin excepciones, se resta la comida habitual del proyecto que cruza su turno.
                registro.comidas.all().delete()
                excepciones = []
                if novedad_codigo == 'asistencia_normal':
                    for tipo, rango in datos_comidas.get(str(emp_pk), {}).items():
                        horas = parsear_rango_horas(rango)
                        if horas:
                            RegistroComida.objects.create(
                                registro=registro,
                                tipo_de_alimentacion=tipo.capitalize(),
                                hora_de_inicio=horas[0],
                                hora_fin=horas[1],
                            )
                            excepciones.append(horas)
                registro.horas_comida = horas_comida_registro(
                    proyecto_actual, maestro_actual.fecha, hora_o_none(h_inicio), hora_o_none(h_fin), excepciones)
                registro.save(update_fields=['horas_comida'])

        # 5. Cambiar el estado del Maestro a "Completado"
        maestro_actual.estado = 'Completado'
        maestro_actual.save()

        # 6. Destruir los datos temporales de la mochila
        for k in ['datos_paso1', 'datos_paso2', 'datos_paso3', 'datos_paso4', 'datos_comidas', 'datos_multiples', 'datos_comentarios', 'datos_historicas', 'fecha_asistencia']:
            if k in request.session:
                del request.session[k]
                
        return redirect(url_quincena_de_fecha(id_proyecto, fecha_str))

    # PROCESO DE VISUALIZACIÓN (MÉTODO GET)
    ids_empleados = list(datos_paso1.keys())
    empleados_db = Empleado.objects.filter(pk__in=ids_empleados).order_by('nombre_completo')

    consolidado = []
    for emp in empleados_db:
        emp_id = str(emp.pk)
        cod_nov = datos_paso1.get(emp_id, '')
        trabaja = cod_nov == 'asistencia_normal'

        # Comida: las excepciones de la herramienta, o la hora habitual del proyecto
        excepciones = [
            f"{tipo.capitalize()}: {texto_hora(rango.split('-')[0])} - {texto_hora(rango.split('-')[1])}"
            for tipo, rango in datos_comidas.get(emp_id, {}).items()
        ] if trabaja else []
        if not trabaja:
            comida = []
        elif excepciones:
            comida = excepciones
        else:
            habitual = horas_comida_habituales(
                proyecto_actual, datetime.strptime(fecha_str, '%Y-%m-%d').date(),
                hora_o_none(convertir_a_24h(datos_paso2.get(emp_id, ''))), hora_o_none(convertir_a_24h(datos_paso3.get(emp_id, ''))),
            )
            comida = [f"{habitual} h (habitual)"]

        consolidado.append({
            'empleado': emp,
            'cod_novedad': cod_nov,
            'hora_entrada_24h': convertir_a_24h(datos_paso2.get(emp_id, '')),
            'hora_salida_24h': convertir_a_24h(datos_paso3.get(emp_id, '')),
            'actividad': datos_paso4.get(emp_id, ''),
            'comida': comida,
            'multiples': [
                f"Salida {texto_hora(r.split('-')[0])} / Ingreso {texto_hora(r.split('-')[1])}"
                for r in datos_multiples.get(emp_id, [])
            ] if trabaja else [],
            'pendientes': [f"{h['tipo']}: {h['cantidad']} h" for h in datos_historicas.get('pendientes', {}).get(emp_id, [])],
            'deducir': [f"{h['tipo']}: {h['cantidad']} h" for h in datos_historicas.get('deducir', {}).get(emp_id, [])],
            'comentario': datos_comentarios.get(emp_id, {}).get('comentario', ''),
            'link': datos_comentarios.get(emp_id, {}).get('link', ''),
        })

    context = {
        'proyecto_actual': proyecto_actual,
        'fecha_seleccionada': fecha_str,
        'consolidado': consolidado,
        'novedades_lista': NOMBRES_NOVEDADES, 
        'actividades_disponibles': actividades_disponibles, 
        'paso_actual': 5,
        'total_pasos': 5,
        'titulo_paso': 'Vista previa y consolidación final del registro',
        'url_atras': f"{reverse('paso_4_actividades', args=[id_proyecto])}?fecha={fecha_str}",
        'errores': errores,
    }
    return render(request, 'registro_asistencia/vista_previa.html', context)


# ---------- Actualizar el registro de un empleado ----------
NOVEDADES_VALIDAS = {valor for valor, _ in RegistroEntradaSalida.NOVEDADES}


def validar_actualizacion(post, novedad_anterior=None):
    """Valida lo que llega del formulario de actualización. Devuelve (datos, errores); si hay errores no se guarda nada.

    Las mismas reglas del registro: horas solo en asistencia, comidas y salidas múltiples dentro del turno,
    hasta 4 horas históricas por pestaña sin repetir el tipo, enlace de soporte con URL válida.
    """
    errores = []
    novedad = post.get('novedad', '')
    if novedad not in NOVEDADES_VALIDAS:
        errores.append("La novedad seleccionada no es válida.")
    asistencia = novedad == 'Asistencia'

    entrada = hora_o_none(post.get('hora_inicio', '')) if asistencia else None
    salida = hora_o_none(post.get('hora_fin', '')) if asistencia else None
    if asistencia:
        if not entrada or not salida:
            errores.append("La hora de ingreso y la de salida son obligatorias en una asistencia.")
        elif entrada == salida:
            errores.append("La hora de ingreso y la de salida no pueden ser iguales.")

    actividad = post.get('actividad', '').strip() if asistencia else ''
    if len(actividad) > 255:
        errores.append("La actividad es demasiado larga.")
    if asistencia and novedad_anterior and novedad_anterior != ASISTENCIA and not actividad:
        errores.append("Al pasar a Asistencia es obligatorio asignar la actividad.")

    comentario = post.get('comentario', '').strip()
    link = post.get('link_del_comentario', '').strip()
    if link:
        try:
            URLValidator()(link)
        except ValidationError:
            errores.append("El enlace de soporte no es una URL válida.")

    # Comidas: solo en asistencia. Las que tienen horario se descuentan; las que se dejan vacías no se tomaron
    comidas = {}
    if asistencia:
        for tipo in TIPOS_COMIDA:
            ini, fin = hora_o_none(post.get(f'comida_{tipo}_inicio', '')), hora_o_none(post.get(f'comida_{tipo}_fin', ''))
            if ini is None and fin is None:
                continue
            if ini is not None and fin is not None and fin > ini:
                comidas[tipo] = (ini, fin)
            else:
                errores.append(f"La comida «{tipo}» necesita una hora de inicio y una de fin posterior (o déjala vacía si no la tomó).")

    # Entradas y salidas múltiples: dentro del turno, después de la entrada y antes de la salida
    multiples = []
    if asistencia:
        for salida_m, ingreso_m in zip(post.getlist('mult_salida'), post.getlist('mult_ingreso')):
            if not salida_m.strip() and not ingreso_m.strip():
                continue
            horas = parsear_rango_libre(f"{salida_m.strip()}-{ingreso_m.strip()}")
            if horas and multiple_dentro_del_turno(entrada, salida, *horas):
                multiples.append(horas)
            else:
                errores.append("Cada salida múltiple necesita salida e ingreso (el ingreso después de la salida), dentro del turno.")
                break
        if not errores and multiples_solapadas(multiples):
            errores.append("Las salidas múltiples no pueden cruzarse entre sí.")

    # Horas históricas: pendientes (a favor) y a deducir (en contra)
    historicas = []
    tipos_validos = {valor for valor, _ in HoraHistorica.TIPO_HORAS}
    for prefijo, accion, nombre in (('pend', 'Pendiente', 'a favor'), ('ded', 'Deducir', 'a deducir')):
        vistos = set()
        filas = 0
        for tipo, cantidad in zip(post.getlist(f'{prefijo}_tipo'), post.getlist(f'{prefijo}_cantidad')):
            if not tipo.strip() and not cantidad.strip():
                continue
            filas += 1
            valor = a_decimal(cantidad)
            if tipo not in tipos_validos or valor is None or not 0 < valor <= 24:
                errores.append(f"Horas históricas {nombre}: elige el tipo y una cantidad mayor que 0 y de máximo 24 h.")
                break
            if tipo in vistos:
                errores.append(f"Horas históricas {nombre}: el tipo «{tipo}» está repetido.")
                break
            vistos.add(tipo)
            historicas.append((accion, tipo, valor.quantize(Decimal('0.01'))))
        if filas > MAX_HORAS_HISTORICAS:
            errores.append(f"Horas históricas {nombre}: máximo {MAX_HORAS_HISTORICAS} filas.")

    datos = {
        'novedad': novedad, 'entrada': entrada, 'salida': salida, 'actividad': actividad or None,
        'comentario': comentario or None, 'link': link or None,
        'comidas': comidas, 'multiples': multiples, 'historicas': historicas,
    }
    return datos, errores


def aplicar_actualizacion(registro, datos):
    """Guarda los cambios del registro y recalcula lo que genera el sistema (jornada base y hora de comida).

    Las horas extra no se guardan: se calculan a partir de estos datos en la tabla del día y en el informe.
    """
    fecha = registro.maestro.fecha
    registro.novedad = datos['novedad']
    registro.hora_inicio, registro.hora_fin = datos['entrada'], datos['salida']
    registro.actividad = datos['actividad']
    registro.comentario, registro.link_del_comentario = datos['comentario'], datos['link']
    registro.horas_base = horas_base_del_dia(fecha)
    # Las comidas escritas son las que se tomaron: no se vuelve a la habitual (una comida vacía no se descuenta)
    registro.horas_comida = horas_de_comidas(datos['comidas'].values())
    registro.save()

    registro.comidas.all().delete()
    for tipo, (ini, fin) in datos['comidas'].items():
        RegistroComida.objects.create(registro=registro, tipo_de_alimentacion=tipo.capitalize(), hora_de_inicio=ini, hora_fin=fin)

    registro.horas_multiples.all().delete()
    for salida_m, ingreso_m in datos['multiples']:
        EntradaSalidaMultiple.objects.create(registro=registro, hora_de_salida_multiple=salida_m, hora_de_ingreso_multiple=ingreso_m)

    registro.horas_historicas.all().delete()
    for accion, tipo, cantidad in datos['historicas']:
        HoraHistorica.objects.create(registro=registro, tipo_de_horas=tipo, accion=accion, cantidad_de_horas=cantidad)


def contexto_campos_registro(proyecto, actividad_actual=None):
    """Lo que necesita componentes/ar_campos_registro.html: novedades, actividades del proyecto, tipos de hora y tope de históricas.

    `actividad_actual` es la actividad ya guardada de un registro: se añade a la lista aunque ya no esté en el proyecto.
    """
    actividades = list(Cumplimiento.objects.filter(proyecto=proyecto).values_list('actividad', flat=True).distinct().order_by('actividad'))
    if actividad_actual and actividad_actual not in actividades:
        actividades.append(actividad_actual)
    return {
        'novedades': RegistroEntradaSalida.NOVEDADES,
        'actividades': actividades,
        'tipos_horas': HoraHistorica.TIPO_HORAS,
        'max_historicas': MAX_HORAS_HISTORICAS,
    }


def respuesta_comidas_habituales(proyecto, fecha, entrada, salida):
    """JSON con las comidas habituales que cruza un turno (los formularios las rellenan al cambiar la hora de ingreso o de salida)."""
    cruces = comidas_habituales_del_turno(proyecto.nombre, fecha, hora_o_none(entrada), hora_o_none(salida))
    return JsonResponse({'comidas': {tipo: [f"{a:%H:%M}", f"{b:%H:%M}"] for tipo, a, b in cruces}})


@permiso_registro('actualizar', json=True)
def actualizar_registro(request, id_proyecto, id_registro):
    """Actualiza el registro de UN empleado en un día.

    GET devuelve el modal con los datos actuales. POST valida y guarda (responde JSON). No se editan los datos
    personales ni lo que calcula el sistema; al guardar se recalculan la jornada, la comida y, con ellas, las horas extra.
    """
    proyecto = get_object_or_404(Proyecto, pk=id_proyecto)
    consulta = RegistroEntradaSalida.objects.select_related('empleado', 'maestro').filter(maestro__proyecto=proyecto)

    if request.method == 'POST':
        with transaction.atomic():
            # El bloqueo hace que dos actualizaciones del mismo registro se atiendan una a la vez
            registro = get_object_or_404(consulta.select_for_update(), pk=id_registro)
            datos, errores = validar_actualizacion(request.POST, registro.novedad)
            if errores:
                return JsonResponse({'success': False, 'errores': errores}, status=400)
            aplicar_actualizacion(registro, datos)
        return JsonResponse({'success': True})

    registro = get_object_or_404(consulta, pk=id_registro)
    fecha = registro.maestro.fecha

    if request.GET.get('habituales'):
        return respuesta_comidas_habituales(proyecto, fecha, request.GET.get('entrada', ''), request.GET.get('salida', ''))

    # Comidas que se muestran: las guardadas; si no hay ninguna y se descontó comida, la habitual que cruza su turno
    guardadas = {c.tipo_de_alimentacion.lower(): (c.hora_de_inicio, c.hora_fin) for c in registro.comidas.all()}
    if not guardadas and registro.novedad == ASISTENCIA and registro.horas_comida > 0:
        guardadas = {tipo: (a, b) for tipo, a, b in comidas_habituales_del_turno(
            proyecto.nombre, fecha, registro.hora_inicio, registro.hora_fin)}
    context = {
        'proyecto_actual': proyecto,
        'registro': registro,
        **contexto_campos_registro(proyecto, registro.actividad),
        'comidas': [(tipo, tipo.capitalize(), guardadas.get(tipo)) for tipo in TIPOS_COMIDA],
        'multiples': registro.horas_multiples.all(),
        'pendientes': registro.horas_historicas.filter(accion='Pendiente'),
        'deducir': registro.horas_historicas.filter(accion='Deducir'),
        'url_guardar': reverse('actualizar_registro', args=[id_proyecto, id_registro]),
    }
    return render(request, 'registro_asistencia/modal_actualizar_registro.html', context)


def dia_de_registro(r):
    """El registro de un empleado en un día, en la forma que usa el cálculo de horas."""
    return DiaEmpleado(
        fecha=r.maestro.fecha,
        novedad=r.novedad,
        entrada=r.hora_inicio,
        salida=r.hora_fin,
        comida_min=int(r.horas_comida * 60) if r.novedad == ASISTENCIA else 0,
        multiples=[(m.hora_de_salida_multiple, m.hora_de_ingreso_multiple) for m in r.horas_multiples.all()],
        historicas=[(h.accion, h.tipo_de_horas, int(h.cantidad_de_horas * 60)) for h in r.horas_historicas.all()],
    )


def datos_quincena(id_proyecto, anio, mes, quincena):
    """Proyecto y rango de fechas de la quincena (404 si no existe)."""
    proyecto = get_object_or_404(Proyecto, pk=id_proyecto)
    if not 1 <= mes <= 12 or quincena not in (1, 2):
        raise Http404("Quincena inválida")
    inicio, fin = rango_quincena(anio, mes, quincena)
    return proyecto, inicio, fin


def fechas_incompletas(proyecto, inicio, fin):
    """Días hábiles de la quincena que todavía no tienen un registro guardado."""
    registradas = set(
        MaestroAsistencia.objects.filter(proyecto=proyecto, estado='Completado', fecha__range=(inicio, fin))
        .values_list('fecha', flat=True)
    )
    return [d for d in dias_habiles(inicio, fin) if d not in registradas]


def quincena_incompleta(proyecto, inicio, fin):
    """True si faltan días por registrar (o no se pudo consultar el calendario de festivos)."""
    try:
        return bool(fechas_incompletas(proyecto, inicio, fin))
    except FestivosNoDisponibles:
        return True


def volver_a_carpetas(id_proyecto, anio, mes, quincena, parametro):
    """La página no carga con días sin registrar: se vuelve a las carpetas y allí sale el aviso."""
    return redirect(f"{reverse('registro_dias', args=[id_proyecto, anio, mes, quincena])}?{parametro}=1")


def contexto_quincena(proyecto, id_proyecto, anio, mes, quincena, inicio, fin):
    """Lo que comparten el informe y el resumen: encabezado, migas y enlace para volver."""
    return {
        'proyecto_actual': proyecto,
        'quincena_codigo': f"Q{quincena}",
        'rango': f"{inicio.day} al {fin.day} de {MESES[mes - 1].lower()} de {anio}",
        'migas': construir_migas(proyecto, anio, mes, quincena),
        'url_volver': reverse('registro_dias', args=[id_proyecto, anio, mes, quincena]),
    }


def validar_quincena_completa(id_proyecto, anio, mes, quincena, destino):
    """Revisa que no falten días por registrar antes de abrir `destino` (el informe o el resumen). Responde JSON."""
    proyecto, inicio, fin = datos_quincena(id_proyecto, anio, mes, quincena)
    try:
        faltan = fechas_incompletas(proyecto, inicio, fin)
    except FestivosNoDisponibles:
        return JsonResponse({
            'ok': False,
            'error': 'No se pudo consultar el calendario de festivos y no hay una copia guardada. Intenta de nuevo en unos minutos.',
        }, status=503)
    if faltan:
        return JsonResponse({'ok': False, 'fechas': [d.isoformat() for d in faltan]}, status=409)
    return JsonResponse({'ok': True, 'url': reverse(destino, args=[id_proyecto, anio, mes, quincena])})


@permiso_registro('informe', json=True)
def informe_quincenal_validar(request, id_proyecto, anio, mes, quincena):
    """Revisa que no falten días por registrar antes de abrir el informe (responde JSON)."""
    return validar_quincena_completa(id_proyecto, anio, mes, quincena, 'informe_quincenal')


def texto_resumen_novedades(resumen, fecha_retiro=None, fecha_ingreso=None):
    """'Sanción: 1 · Vacaciones: 3'. El retiro muestra la fecha en que se retiró en vez de la cantidad;
    el empleado nuevo (primer registro después del inicio de la quincena) empieza con su fecha de ingreso."""
    partes = [f"Nuevo ingreso: {fecha_ingreso:%d/%m/%Y}"] if fecha_ingreso else []
    partes += [
        f"Retiro: {fecha_retiro:%d/%m/%Y}" if novedad == 'Retiro' and fecha_retiro else f"{novedad}: {cantidad}"
        for novedad, cantidad in sorted(resumen.items())
    ]
    return ' · '.join(partes) if partes else 'Sin novedades'


def rangos_de_fechas(fechas):
    """['05/10 al 08/10', '14/10']: junta las fechas seguidas; los fines de semana y festivos entre dos fechas no las separan."""
    rangos = []
    for fecha in sorted(fechas):
        if rangos and not any(True for _ in dias_habiles(rangos[-1][1] + timedelta(days=1), fecha - timedelta(days=1))):
            rangos[-1][1] = fecha
        else:
            rangos.append([fecha, fecha])
    return [f"{a:%d/%m}" if a == b else f"{a:%d/%m} al {b:%d/%m}" for a, b in rangos]


def dias_de_descanso(dias, desde, hasta):
    """Sábados, domingos y festivos entre `desde` y `hasta` en los que el empleado no trabajó ni tuvo otra novedad.

    Aunque no tengan registro cuentan como descanso; los que sí trabajó son asistencia, no descanso. Un día registrado como
    «Descanso» fuera de ese rango (por ejemplo un sábado posterior al último día hábil de un retirado) tampoco cuenta:
    ya es de retiro y no se puede contar dos veces.
    """
    ocupados = {d.fecha for d in dias if d.novedad != DESCANSO}
    descansos = {d.fecha for d in dias if d.novedad == DESCANSO and desde <= d.fecha <= hasta}
    dia = desde
    while dia <= hasta:
        if not es_habil(dia) and dia not in ocupados:
            descansos.add(dia)
        dia += timedelta(days=1)
    return len(descansos)


def columnas_novedades(dias, inicio, fin, fecha_ingreso=None, fecha_retiro=None):
    """Lo que el informe muestra por novedad: cuántas asistencias, descansos, ausentismos e incapacidades tuvo y en qué fechas.

    Los descansos se cuentan desde que ingresó el empleado hasta el fin de la quincena o, si se retiró, hasta el último día
    hábil anterior a su retiro: los sábados, domingos y festivos entre ese día y la fecha del retiro ya son del retiro
    (retiro el lunes tras trabajar hasta el viernes: el fin de semana no es descanso).
    """
    fechas = {}
    for dia in dias:
        fechas.setdefault(dia.novedad, []).append(dia.fecha)
    desde_descanso = fecha_ingreso or inicio
    hasta_descanso = ultimo_dia_habil_antes(fecha_retiro, desde_descanso) if fecha_retiro else fin
    return {
        'asistencias': len(fechas.get(ASISTENCIA, [])),
        'descansos': dias_de_descanso(dias, desde_descanso, hasta_descanso) if hasta_descanso else 0,
        'licencias': len(fechas.get('Licencia', [])),
        'sanciones': len(fechas.get('Sanción', [])),
        'vacaciones': len(fechas.get('Vacaciones', [])),
        'permisos_remunerados': len(fechas.get('Permiso remunerado', [])),
        'permisos_no_remunerados': len(fechas.get('Permiso no remunerado', [])),
        'ausentismos': len(fechas.get('Ausentismo', [])),
        'fechas_ausentismo': [f"{f:%d/%m}" for f in sorted(fechas.get('Ausentismo', []))],
        'inc_general': len(fechas.get('Incapacidad general', [])),
        'rangos_inc_general': rangos_de_fechas(fechas.get('Incapacidad general', [])),
        'inc_laboral': len(fechas.get('Incapacidad laboral', [])),
        'rangos_inc_laboral': rangos_de_fechas(fechas.get('Incapacidad laboral', [])),
    }


def filas_del_informe(proyecto, inicio, fin):
    """(filas, totales) del informe de la quincena. Lanza FestivosNoDisponibles si falta el calendario de festivos."""
    registros = RegistroEntradaSalida.objects.filter(
        maestro__proyecto=proyecto, maestro__estado='Completado', maestro__fecha__range=(inicio, fin),
    ).select_related('empleado', 'maestro').prefetch_related('horas_multiples', 'horas_historicas')

    por_empleado = {}
    for r in registros:
        por_empleado.setdefault(r.empleado, []).append(dia_de_registro(r))

    jornadas = jornadas_por_dia()
    filas = []
    for empleado, dias in por_empleado.items():
        res = liquidar(dias, fin, jornadas, inicio)
        filas.append({
            'empleado': empleado,
            'horas_totales': res['total'],
            'heo': res['extras'][HEO],
            'nocturna': res['extras'][NOCTURNA],
            'dom_diurna': res['extras'][DOM_DIURNA],
            'dom_nocturna': res['extras'][DOM_NOCTURNA],
            **columnas_novedades(dias, inicio, fin, res['fecha_ingreso'], res['fecha_retiro']),
            'fecha_retiro': res['fecha_retiro'],
            'fecha_ingreso': res['fecha_ingreso'],
            'dias_desde_ingreso': res['dias_desde_ingreso'],
            'umbral': res['umbral'],
            'horas_base': res['horas_base'],
            'dias_retiro': res['dias_retiro'],
            'horas_deducir': res['horas_deducir'],
            'novedades_deducir': res['novedades_deducir'],
            'cuadra': res['cuadra'],
        })
    filas.sort(key=lambda f: f['empleado'].nombre_completo)

    # Total de cada columna con números: las de horas con dos decimales y las de días como enteros
    totales = {
        clave: sum((f[clave] for f in filas), Decimal('0.00'))
        for clave in ('horas_base', 'horas_totales', 'heo', 'nocturna', 'dom_diurna', 'dom_nocturna', 'horas_deducir')
    }
    totales.update({
        clave: sum(f[clave] for f in filas)
        for clave in ('ausentismos', 'licencias', 'asistencias', 'sanciones', 'descansos', 'vacaciones',
                      'permisos_remunerados', 'permisos_no_remunerados', 'inc_laboral', 'inc_general', 'dias_retiro')
    })
    return filas, totales


def cedulas_con_historial(inicio, empleados):
    """Cédulas (de `empleados`) con algún registro guardado antes de `inicio`, en cualquier proyecto: no son empleados nuevos."""
    return set(RegistroEntradaSalida.objects.filter(
        maestro__estado='Completado', maestro__fecha__lt=inicio, empleado__in=list(empleados),
    ).values_list('empleado_id', flat=True).distinct())


@permiso_registro('informe')
def informe_quincenal(request, id_proyecto, anio, mes, quincena):
    """Informe de la quincena: horas extra por tipo, horas a deducir y novedades de cada empleado."""
    proyecto, inicio, fin = datos_quincena(id_proyecto, anio, mes, quincena)
    if quincena_incompleta(proyecto, inicio, fin):
        return volver_a_carpetas(id_proyecto, anio, mes, quincena, 'informe_incompleto')

    try:
        filas, totales = filas_del_informe(proyecto, inicio, fin)
    except FestivosNoDisponibles:
        return volver_a_carpetas(id_proyecto, anio, mes, quincena, 'informe_incompleto')

    # Solo el empleado nuevo (marcado «Nuevo» y sin registros antes de la quincena) puede cambiar su fecha de ingreso.
    # El servidor lo vuelve a comprobar en cada vista de fecha_ingreso.py; esto solo decide si se enciende el ícono.
    veteranos = cedulas_con_historial(inicio, [f['empleado'] for f in filas if f['fecha_ingreso']])
    for f in filas:
        f['ingreso_editable'] = bool(f['fecha_ingreso']) and f['empleado'].pk not in veteranos

    context = {
        **contexto_quincena(proyecto, id_proyecto, anio, mes, quincena, inicio, fin),
        'filas': filas,
        'totales': totales,
        'url_descarga': reverse('descargar_informe_excel', args=[id_proyecto, anio, mes, quincena]),
        # 'CEDULA' es un comodín que fecha_ingreso.js cambia por la cédula del empleado seleccionado
        'url_fecha_ingreso': reverse('fecha_ingreso', args=[id_proyecto, anio, mes, quincena, 'CEDULA']),
    }
    return render(request, 'registro_asistencia/informe_quincenal.html', context)


# Color del nombre del empleado en el informe (pantalla: informe_quincenal.css): nuevo en verde, retirado en rojo
COLOR_NUEVO = '3D7A12'
COLOR_RETIRADO = 'A40A02'


def _titulo_con_unidad(titulo, unidad):
    """Encabezado de Excel: el título en negrita y, debajo, la unidad en letra pequeña (como en la pantalla)."""
    if not unidad:
        return titulo
    return CellRichText(TextBlock(InlineFont(b=True, color="FFFFFF", sz=11), titulo + "\n"),
                        TextBlock(InlineFont(color="FFFFFF", sz=8), unidad))


def _dar_formato_excel(ws, filas_encabezado, hay_total, anchos):
    """Formato común de los Excel de la quincena: bordes, encabezado verde, total en negrita y primera columna fija."""
    relleno = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
    borde = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    for fila in ws.iter_rows():
        for celda in fila:
            celda.border = borde
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    for fila in ws.iter_rows(min_row=1, max_row=filas_encabezado):
        for celda in fila:
            celda.fill = relleno
            celda.font = Font(color="FFFFFF", bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[fila[0].row].height = 36
    if hay_total:
        for celda in ws[ws.max_row]:
            celda.font = Font(bold=True)
    for indice, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(indice)].width = ancho
    ws.freeze_panes = ws.cell(filas_encabezado + 1, 2)


def _respuesta_excel(wb, nombre):
    nombre = get_valid_filename(nombre.replace(' ', '_'))
    respuesta = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}"'
    wb.save(respuesta)
    return respuesta


def _celda_con_detalle(cantidad, detalle):
    """Número de días; si hay fechas o rangos, el número y debajo el detalle (como se ve en el informe)."""
    return f"{cantidad}\n{detalle}" if cantidad and detalle else cantidad


@permiso_registro('informe')
@permiso_registro('descargar')
def descargar_informe_excel(request, id_proyecto, anio, mes, quincena):
    """Excel del informe quincenal: las mismas columnas y filas que se ven en pantalla, más la fila de total."""
    proyecto, inicio, fin = datos_quincena(id_proyecto, anio, mes, quincena)
    if quincena_incompleta(proyecto, inicio, fin):
        return volver_a_carpetas(id_proyecto, anio, mes, quincena, 'informe_incompleto')
    try:
        filas, totales = filas_del_informe(proyecto, inicio, fin)
    except FestivosNoDisponibles:
        return volver_a_carpetas(id_proyecto, anio, mes, quincena, 'informe_incompleto')

    # Igual que la tabla: (título, unidad) por columna; "Permisos" e "Incapacidades" agrupan dos columnas cada uno
    columnas = [
        ('Nombre del empleado', None), ('Horas base de la quincena', 'horas'), ('Horas laborales totales', 'horas'),
        ('Extra ordinaria diurna', 'horas'),
        ('Extra nocturna ordinaria', 'horas'), ('Extra dominical / festiva diurna', 'horas'),
        ('Extra dominical / festiva nocturna', 'horas'), ('Ausentismo', 'días y fechas'), ('Licencias', 'días'),
        ('Asistencias', 'días'), ('Suspendido', 'días'), ('Descanso', 'días'), ('Vacaciones', 'días'),
        ('Remunerado', None), ('No remunerado', None), ('Accidente de trabajo', None), ('Enfermedad general', None),
        ('Retiro', 'días sin laborar'), ('Horas a deducir', None), ('Novedades a deducir', None),
    ]
    grupos = {14: ('Permisos', 'días'), 16: ('Incapacidades', 'días y rango')}   # primera columna del grupo
    total_columnas = len(columnas)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Informe Q{quincena}"
    ws.append([None] * total_columnas)
    ws.append([None] * total_columnas)
    for numero, (titulo, unidad) in enumerate(columnas, start=1):
        if numero in grupos:
            ws.cell(1, numero, _titulo_con_unidad(*grupos[numero]))
            ws.merge_cells(start_row=1, start_column=numero, end_row=1, end_column=numero + 1)
            ws.cell(2, numero, _titulo_con_unidad(titulo, unidad))
        elif numero - 1 in grupos:
            ws.cell(2, numero, _titulo_con_unidad(titulo, unidad))
        else:
            ws.cell(1, numero, _titulo_con_unidad(titulo, unidad))
            ws.merge_cells(start_row=1, start_column=numero, end_row=2, end_column=numero)

    for f in filas:
        ws.append([
            texto_excel(f['empleado'].nombre_completo),
            float(f['horas_base']), float(f['horas_totales']), float(f['heo']), float(f['nocturna']),
            float(f['dom_diurna']), float(f['dom_nocturna']),
            _celda_con_detalle(f['ausentismos'], ', '.join(f['fechas_ausentismo'])),
            f['licencias'], f['asistencias'], f['sanciones'], f['descansos'], f['vacaciones'],
            f['permisos_remunerados'], f['permisos_no_remunerados'],
            _celda_con_detalle(f['inc_laboral'], '\n'.join(f['rangos_inc_laboral'])),
            _celda_con_detalle(f['inc_general'], '\n'.join(f['rangos_inc_general'])),
            _celda_con_detalle(f['dias_retiro'], f"Retiro: {f['fecha_retiro']:%d/%m/%Y}" if f['fecha_retiro'] else ''),
            float(f['horas_deducir']),
            texto_excel('\n'.join(f['novedades_deducir'])) or '--',
        ])
    if filas:
        # Un total por cada columna con números, en el mismo orden de la tabla; «Novedades a deducir» es texto y no suma
        claves_total = ('horas_base', 'horas_totales', 'heo', 'nocturna', 'dom_diurna', 'dom_nocturna',
                        'ausentismos', 'licencias', 'asistencias', 'sanciones', 'descansos', 'vacaciones',
                        'permisos_remunerados', 'permisos_no_remunerados', 'inc_laboral', 'inc_general', 'dias_retiro',
                        'horas_deducir')
        ws.append(['Total', *(float(totales[c]) if isinstance(totales[c], Decimal) else totales[c] for c in claves_total), None])
    else:
        ws.append(['Esta quincena no tiene registros guardados.'])

    # Igual que en pantalla: el nombre en verde si es un empleado nuevo y en rojo si se retiró
    for fila_excel, f in enumerate(filas, start=3):
        if f['fecha_retiro']:
            ws.cell(fila_excel, 1).font = Font(color=COLOR_RETIRADO)
        elif f['fecha_ingreso']:
            ws.cell(fila_excel, 1).font = Font(color=COLOR_NUEVO)

    anchos = [38] + [18] * (total_columnas - 2) + [60]
    _dar_formato_excel(ws, 2, bool(filas), anchos)
    return _respuesta_excel(wb, f"informe_{proyecto.nombre}_{anio}-{mes:02d}_Q{quincena}.xlsx")


# ---------- Resumen de actividades de la quincena ----------
# Perfiles del resumen, en el orden de las columnas, y el cargo (en minúsculas) de cada empleado que cae en cada uno.
# Se cambia aquí: «Oficial senior» y «Oficial junior» suman juntos en «Oficial»; «Ayudante entendido» cuenta como
# ayudante senior y «Ayudante» como ayudante junior.
PERFILES_RESUMEN = ['Maestro', 'Almacenista', 'Oficial', 'Ayudante senior', 'Ayudante junior']
CARGO_A_PERFIL = {
    'maestro de obra': 'Maestro', 'almacenista': 'Almacenista',
    'oficial senior': 'Oficial', 'oficial junior': 'Oficial',
    'ayudante entendido': 'Ayudante senior', 'ayudante': 'Ayudante junior',
}
PERFIL_OTROS = 'Otros cargos'
SIN_ACTIVIDAD = 'Sin actividad asignada'


def perfil_de_cargo(cargo):
    return CARGO_A_PERFIL.get((cargo or '').strip().lower(), PERFIL_OTROS)


def resumen_de_actividades(registros):
    """Horas efectivas trabajadas por actividad y por perfil: [(actividad, {perfil: minutos})] y los minutos por perfil.

    Solo cuentan los días de asistencia: horas = (salida - entrada) - comida - salidas temporales, como en el informe.
    """
    por_actividad = {}
    for r in registros:
        if r.novedad != ASISTENCIA:
            continue
        minutos = minutos_efectivos(dia_de_registro(r))
        if minutos <= 0:
            continue
        por_perfil = por_actividad.setdefault((r.actividad or '').strip() or SIN_ACTIVIDAD, {})
        perfil = perfil_de_cargo(r.empleado.cargo)
        por_perfil[perfil] = por_perfil.get(perfil, 0) + minutos
    # Orden alfabético; «Sin actividad asignada» siempre al final
    orden = sorted(por_actividad, key=lambda a: (a == SIN_ACTIVIDAD, a.lower()))
    return [(a, por_actividad[a]) for a in orden]


@permiso_registro('resumen', json=True)
def resumen_actividades_validar(request, id_proyecto, anio, mes, quincena):
    """Igual que el informe: el resumen no se abre si faltan días por registrar (responde JSON)."""
    return validar_quincena_completa(id_proyecto, anio, mes, quincena, 'resumen_actividades')


def filas_del_resumen(proyecto, inicio, fin):
    """(perfiles, filas, totales) del resumen de actividades de la quincena."""
    registros = RegistroEntradaSalida.objects.filter(
        maestro__proyecto=proyecto, maestro__estado='Completado', maestro__fecha__range=(inicio, fin),
    ).select_related('empleado', 'maestro').prefetch_related('horas_multiples')
    resumen = resumen_de_actividades(registros)

    perfiles = list(PERFILES_RESUMEN)
    if any(PERFIL_OTROS in por_perfil for _, por_perfil in resumen):
        perfiles.append(PERFIL_OTROS)               # horas de cargos que no son de ninguna columna (p. ej. administrativos)

    filas = [{
        'actividad': actividad,
        'celdas': [horas_de(por_perfil.get(p, 0)) for p in perfiles],
        'total': horas_de(sum(por_perfil.values())),
    } for actividad, por_perfil in resumen]
    totales = {
        'celdas': [horas_de(sum(por_perfil.get(p, 0) for _, por_perfil in resumen)) for p in perfiles],
        'total': horas_de(sum(sum(por_perfil.values()) for _, por_perfil in resumen)),
    }
    return perfiles, filas, totales


@permiso_registro('resumen')
def resumen_actividades(request, id_proyecto, anio, mes, quincena):
    """Tabla con las horas que se gastaron en cada actividad durante la quincena, separadas por perfil."""
    proyecto, inicio, fin = datos_quincena(id_proyecto, anio, mes, quincena)
    if quincena_incompleta(proyecto, inicio, fin):
        return volver_a_carpetas(id_proyecto, anio, mes, quincena, 'resumen_incompleto')

    perfiles, filas, totales = filas_del_resumen(proyecto, inicio, fin)
    context = {
        **contexto_quincena(proyecto, id_proyecto, anio, mes, quincena, inicio, fin),
        'perfiles': perfiles,
        'filas': filas,
        'totales': totales,
        'url_descarga': reverse('descargar_resumen_excel', args=[id_proyecto, anio, mes, quincena]),
    }
    return render(request, 'registro_asistencia/resumen_actividades.html', context)


@permiso_registro('resumen')
@permiso_registro('descargar')
def descargar_resumen_excel(request, id_proyecto, anio, mes, quincena):
    """Excel del resumen de actividades: las mismas columnas y filas que se ven en pantalla, más la fila de total."""
    proyecto, inicio, fin = datos_quincena(id_proyecto, anio, mes, quincena)
    if quincena_incompleta(proyecto, inicio, fin):
        return volver_a_carpetas(id_proyecto, anio, mes, quincena, 'resumen_incompleto')

    perfiles, filas, totales = filas_del_resumen(proyecto, inicio, fin)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Resumen Q{quincena}"
    ws.append([_titulo_con_unidad('Actividad', None)] + [_titulo_con_unidad(p, 'horas') for p in perfiles]
              + [_titulo_con_unidad('Total', 'horas')])
    for f in filas:
        ws.append([texto_excel(f['actividad']), *(float(h) for h in f['celdas']), float(f['total'])])
    if filas:
        ws.append(['Total', *(float(h) for h in totales['celdas']), float(totales['total'])])
    else:
        ws.append(['Esta quincena no tiene horas trabajadas registradas.'])

    anchos = [60] + [16] * (len(perfiles) + 1)
    _dar_formato_excel(ws, 1, bool(filas), anchos)
    return _respuesta_excel(wb, f"resumen_actividades_{proyecto.nombre}_{anio}-{mes:02d}_Q{quincena}.xlsx")


@permiso_registro('eliminar', json=True)
def eliminar_dias(request, id_proyecto):
    """Elimina los días (maestro + sus registros) seleccionados en la vista de carpetas."""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Método no permitido'}, status=405)
    proyecto_actual = get_object_or_404(Proyecto, pk=id_proyecto)
    try:
        fechas = [datetime.strptime(f, "%Y-%m-%d").date() for f in json.loads(request.body).get('fechas', [])]
    except (ValueError, TypeError, AttributeError):
        return JsonResponse({'success': False, 'error': 'Datos inválidos'}, status=400)
    if not fechas:
        return JsonResponse({'success': False, 'error': 'No se seleccionó ningún día'}, status=400)

    # El delete en cascada borra también los RegistroEntradaSalida de cada día.
    MaestroAsistencia.objects.filter(proyecto=proyecto_actual, fecha__in=fechas).delete()
    return JsonResponse({'success': True, 'eliminados': len(fechas)})


@permiso_registro('descargar')
def descargar_dias_excel(request, id_proyecto, anio, mes, quincena):
    """Excel de la quincena: una hoja por cada día guardado."""
    proyecto_actual = get_object_or_404(Proyecto, pk=id_proyecto)
    if not 1 <= mes <= 12 or quincena not in (1, 2):
        raise Http404("Quincena inválida")

    inicio, fin = rango_quincena(anio, mes, quincena)
    maestros = MaestroAsistencia.objects.filter(
        proyecto=proyecto_actual, estado='Completado', fecha__range=(inicio, fin),
    ).order_by('fecha')

    encabezados = ['Código', 'Fecha', 'Nombre del empleado', 'Hora de ingreso', 'Hora de salida',
                    'Tiempo de comida', 'Actividad', 'Novedad']
    relleno = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
    borde = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for maestro in maestros:
        ws = wb.create_sheet(title=maestro.fecha.strftime("%d-%m-%Y"))
        ws.append(encabezados)
        registros = maestro.registros_empleados.select_related('empleado').order_by('empleado__nombre_completo')
        for r in registros:
            ws.append([
                texto_excel(r.empleado.cedula) or r.empleado.pk,
                maestro.fecha.strftime("%d/%m/%Y"),
                texto_excel(r.empleado.nombre_completo),
                r.hora_inicio.strftime("%I:%M %p") if r.hora_inicio else '--:--',
                r.hora_fin.strftime("%I:%M %p") if r.hora_fin else '--:--',
                '--',
                texto_excel(r.actividad) or 'No asignada',
                texto_excel(r.novedad),
            ])
        for fila in ws.iter_rows():
            for celda in fila:
                celda.border = borde
                celda.alignment = Alignment(vertical="top", wrap_text=True)
        for celda in ws[1]:
            celda.fill = relleno
            celda.font = Font(color="FFFFFF", bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for letra, ancho in zip("ABCDEFGH", (16, 12, 34, 16, 16, 16, 50, 22)):
            ws.column_dimensions[letra].width = ancho

    if not wb.sheetnames:
        wb.create_sheet(title="Sin registros").append(['Esta quincena no tiene días guardados.'])

    nombre = get_valid_filename(f"asistencia_{proyecto_actual.nombre}_{anio}-{mes:02d}_Q{quincena}.xlsx".replace(' ', '_'))
    respuesta = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}"'
    wb.save(respuesta)
    return respuesta
